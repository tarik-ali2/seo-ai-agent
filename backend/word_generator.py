import json, os, zipfile
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# ── Color palette ─────────────────────────────────────────────────────────────
C_CRITICAL = RGBColor(0xD3, 0x00, 0x00)   # Red
C_HIGH     = RGBColor(0xE6, 0x74, 0x00)   # Orange
C_MEDIUM   = RGBColor(0xC9, 0xA7, 0x00)   # Yellow-dark
C_LOW      = RGBColor(0x18, 0x80, 0x38)   # Green
C_BLUE     = RGBColor(0x1A, 0x73, 0xE8)   # Brand blue
C_DARK     = RGBColor(0x20, 0x20, 0x20)
C_GRAY     = RGBColor(0x66, 0x66, 0x66)
C_WHITE    = RGBColor(0xFF, 0xFF, 0xFF)

PRIORITY_COLOR = {"critical": C_CRITICAL, "high": C_HIGH, "medium": C_MEDIUM, "low": C_LOW}
PRIORITY_FILL  = {"critical": "FFE5E5", "high": "FFF0E0", "medium": "FFFBE0", "low": "E8F5E9"}
PRIORITY_LABEL = {"critical": "🔴 CRITICAL", "high": "🟠 HIGH", "medium": "🟡 MEDIUM", "low": "🟢 LOW"}


# ── XML / cell helpers ────────────────────────────────────────────────────────

def _set_cell_bg(cell, hex_color: str):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_borders(cell, color="CCCCCC", size="4"):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for side in ["top", "left", "bottom", "right"]:
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)
        tcBorders.append(el)
    tcPr.append(tcBorders)


def _cell_text(cell, text: str, bold=False, color=None, size=9, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run(str(text) if text else "—")
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def _add_shading(paragraph, hex_color: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


# ── Document helpers ──────────────────────────────────────────────────────────

def _h1(doc, text: str):
    p = doc.add_heading("", level=1)
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.bold = True
    run.font.color.rgb = C_BLUE
    return p


def _h2(doc, text: str):
    p = doc.add_heading("", level=2)
    p.clear()
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.bold = True
    run.font.color.rgb = C_DARK
    return p


def _h3(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.bold = True
    run.font.color.rgb = C_DARK
    return p


def _kv(doc, key: str, value, highlight=False):
    p = doc.add_paragraph()
    k = p.add_run(f"{key}: ")
    k.bold = True
    k.font.size = Pt(10)
    v = p.add_run(str(value) if value else "— not set —")
    v.font.size = Pt(10)
    if highlight:
        v.font.color.rgb = C_CRITICAL if not value else C_LOW
    return p


def _code(doc, code: str, label: str = ""):
    if label:
        lp = doc.add_paragraph()
        lr = lp.add_run(f"  {label}")
        lr.font.size = Pt(8)
        lr.font.color.rgb = C_GRAY
        lr.font.italic = True

    p = doc.add_paragraph()
    _add_shading(p, "1E1E1E")
    run = p.add_run(code.strip())
    run.font.name = "Courier New"
    run.font.size = Pt(8.5)
    run.font.color.rgb = C_WHITE
    return p


def _divider(doc):
    p = doc.add_paragraph()
    p.add_run("─" * 90)
    p.runs[0].font.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
    p.runs[0].font.size = Pt(7)


def _page_break(doc):
    doc.add_page_break()


def _footer(doc, url: str):
    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        p.text = f"SEO AI Agent Report  |  {url}  |  Generated: {datetime.now().strftime('%d %b %Y %H:%M')}"
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.color.rgb = C_GRAY


# ── Score badge paragraph ─────────────────────────────────────────────────────

def _score_badge(doc, label: str, score: int):
    p = doc.add_paragraph()
    r = p.add_run(f"  {label}: ")
    r.font.bold = True
    r.font.size = Pt(11)
    score_run = p.add_run(f"{score}/100  ")
    score_run.font.size = Pt(13)
    score_run.font.bold = True
    if score >= 80:
        score_run.font.color.rgb = C_LOW
    elif score >= 50:
        score_run.font.color.rgb = C_MEDIUM
    else:
        score_run.font.color.rgb = C_CRITICAL
    return p


# ── Priority issues table ─────────────────────────────────────────────────────

def _issues_table(doc, issues: list):
    if not issues:
        p = doc.add_paragraph("✅ No issues found on this page.")
        p.runs[0].font.color.rgb = C_LOW
        return

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT

    # Set column widths
    widths = [Cm(2.5), Cm(3.0), Cm(6.5), Cm(6.5)]
    for i, col in enumerate(table.columns):
        for cell in col.cells:
            cell.width = widths[i]

    # Header row
    hdr = table.rows[0].cells
    for cell, text in zip(hdr, ["Priority", "Element", "Problem", "Fix"]):
        _set_cell_bg(cell, "1A73E8")
        _cell_text(cell, text, bold=True, color=C_WHITE, size=9)

    for issue in issues:
        pri = issue.get("priority", "low")
        row = table.add_row().cells
        _set_cell_bg(row[0], PRIORITY_FILL.get(pri, "FFFFFF"))
        _set_cell_bg(row[1], "FAFAFA")
        _set_cell_bg(row[2], "FAFAFA")
        _set_cell_bg(row[3], "F0FFF4")
        _cell_text(row[0], PRIORITY_LABEL.get(pri, pri), bold=True, color=PRIORITY_COLOR.get(pri, C_DARK), size=8)
        _cell_text(row[1], issue.get("element", ""), size=9)
        _cell_text(row[2], issue.get("problem", ""), size=9)
        _cell_text(row[3], issue.get("fix", ""), size=9, color=RGBColor(0x0D, 0x6E, 0x3B))
        for c in row:
            _set_cell_borders(c)


# ── Before / After comparison table ──────────────────────────────────────────

def _before_after_table(doc, rows: list):
    """rows = [(element, current, suggested), ...]"""
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"

    hdr = table.rows[0].cells
    for cell, text, bg in zip(hdr, ["SEO Element", "Current (Problems)", "Suggested (Copy This)"],
                               ["2C3E50", "E74C3C", "27AE60"]):
        _set_cell_bg(cell, bg)
        _cell_text(cell, text, bold=True, color=C_WHITE, size=9)

    for element, current, suggested in rows:
        row = table.add_row().cells
        _set_cell_bg(row[0], "EAF0FB")
        _set_cell_bg(row[1], "FFF5F5")
        _set_cell_bg(row[2], "F0FFF4")
        _cell_text(row[0], element, bold=True, size=9)
        current_str = str(current) if current else "❌ MISSING"
        _cell_text(row[1], current_str[:200], size=8.5,
                   color=C_CRITICAL if not current else C_DARK)
        _cell_text(row[2], str(suggested)[:200] if suggested else "—", size=8.5, color=RGBColor(0x0D, 0x6E, 0x3B))
        for c in row:
            _set_cell_borders(c)


# ── Pages overview table ──────────────────────────────────────────────────────

def _pages_overview_table(doc, pages_analysis: list):
    """Table: #, URL, Page Title, Short Summary (meta or body snippet)."""
    if not pages_analysis:
        doc.add_paragraph("No pages crawled.")
        return

    _kv(doc, "Total Pages Crawled", len(pages_analysis))
    doc.add_paragraph()

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"

    headers = ["#", "Page URL", "Title", "Summary (Meta / Content)"]
    bg_colors = ["1A73E8", "1A73E8", "1A73E8", "1A73E8"]
    for cell, text, bg in zip(table.rows[0].cells, headers, bg_colors):
        _set_cell_bg(cell, bg)
        _cell_text(cell, text, bold=True, color=C_WHITE, size=9)

    # column widths
    col_widths = [Cm(1.0), Cm(5.5), Cm(4.5), Cm(7.5)]
    for i, col in enumerate(table.columns):
        for cell in col.cells:
            cell.width = col_widths[i]

    for idx, page in enumerate(pages_analysis, 1):
        url = page.get("url", "")
        title = page.get("current", {}).get("title") or page.get("title", "") or "—"
        meta = (
            page.get("current", {}).get("meta_description")
            or page.get("meta_description", "")
            or page.get("body_text_sample", "")[:150]
            or "—"
        )
        # trim long strings
        url_display = url if len(url) <= 70 else "…" + url[-67:]
        title_display = title[:80]
        summary_display = meta[:180]

        row = table.add_row().cells
        bg = "F8F9FA" if idx % 2 == 0 else "FFFFFF"
        for c in row:
            _set_cell_bg(c, bg)
            _set_cell_borders(c)

        _cell_text(row[0], str(idx), bold=True, size=9,
                   align=WD_ALIGN_PARAGRAPH.CENTER,
                   color=C_BLUE)
        _cell_text(row[1], url_display, size=8, color=C_GRAY)
        _cell_text(row[2], title_display, size=9, color=C_DARK)
        _cell_text(row[3], summary_display, size=8.5, color=RGBColor(0x44, 0x44, 0x44))


# ── Keyword table ─────────────────────────────────────────────────────────────

def _keyword_table(doc, keywords: list):
    if not keywords:
        doc.add_paragraph("No keywords extracted — page has too little content.")
        return
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for c, t in zip(hdr, ["Rank", "Keyword / Phrase", "Frequency"]):
        _set_cell_bg(c, "1A73E8")
        _cell_text(c, t, bold=True, color=C_WHITE, size=9)
    for i, (kw, freq) in enumerate(keywords[:15], 1):
        row = table.add_row().cells
        bg = "EAF0FB" if i % 2 == 0 else "FFFFFF"
        for c in row:
            _set_cell_bg(c, bg)
            _set_cell_borders(c)
        _cell_text(row[0], f"#{i}", bold=(i <= 3), size=9)
        _cell_text(row[1], kw, bold=(i <= 3), size=9, color=C_BLUE if i <= 3 else C_DARK)
        _cell_text(row[2], str(freq), size=9)


# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 1 — Technical SEO Audit
# ══════════════════════════════════════════════════════════════════════════════

def create_technical_seo_doc(technical_data: dict, output_path: str, pages_analysis: list = None):
    doc = Document()
    _setup_margins(doc)

    # Cover
    _cover(doc, "Technical SEO Audit Report", technical_data.get("base_url", ""),
           technical_data.get("score", 0))

    # Summary cards
    _h1(doc, "1. Overview")
    summary = technical_data.get("summary", {})
    summary_rows = [
        ("Website", technical_data.get("base_url", "")),
        ("Audit Date", datetime.now().strftime("%d %B %Y")),
        ("HTTPS", "✅ Enabled" if technical_data.get("https_enabled") else "❌ NOT enabled — CRITICAL"),
        ("robots.txt", "✅ Found" if technical_data.get("robots_txt_found") else "❌ Not found"),
        ("XML Sitemap", f"✅ {technical_data['sitemap'].get('url_count',0)} URLs" if technical_data.get("sitemap",{}).get("found") else "❌ Not found"),
        ("Pages Crawled", str(technical_data.get("total_pages_crawled", 0))),
        ("Critical Issues", str(summary.get("critical", 0))),
        ("High Priority Issues", str(summary.get("high", 0))),
    ]
    for k, v in summary_rows:
        _kv(doc, k, v)

    _score_badge(doc, "Overall Technical Score", technical_data.get("score", 0))
    _divider(doc)

    # Issues table
    _h1(doc, "2. Issues Found (Priority Order)")
    issues = technical_data.get("issues", [])
    _issues_table(doc, issues)
    doc.add_paragraph()

    # Crawled Pages Overview
    _page_break(doc)
    _h1(doc, "3. Crawled Pages Overview")
    doc.add_paragraph(
        "All pages discovered and crawled during this audit. "
        "Summary column shows meta description or page content snippet."
    )
    doc.add_paragraph()
    _pages_overview_table(doc, pages_analysis or [])

    # Recommendations
    recs = technical_data.get("recommendations", [])
    if recs:
        _h1(doc, "4. What's Working Well")
        for r in recs:
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(f"✅ {r}")
            run.font.color.rgb = C_LOW
            run.font.size = Pt(10)

    _page_break(doc)

    # robots.txt
    _h1(doc, "4. robots.txt Analysis")
    robots = technical_data.get("robots_txt_content", "")
    if robots and "not found" not in robots.lower():
        _code(doc, robots[:1500], "Current robots.txt")
    else:
        doc.add_paragraph("❌ robots.txt not found. Recommended content:")
        _code(doc, """User-agent: *
Allow: /
Disallow: /admin/
Disallow: /cart/
Disallow: /checkout/
Disallow: /?*
Disallow: /user/
Disallow: /api/

Sitemap: https://yoursite.com/sitemap.xml""", "Recommended robots.txt")

    # Sitemap
    _h1(doc, "5. XML Sitemap")
    sitemap = technical_data.get("sitemap", {})
    if sitemap.get("found"):
        _kv(doc, "Sitemap URL", sitemap.get("url", ""))
        _kv(doc, "Total URLs", sitemap.get("url_count", 0))
        _kv(doc, "Status", "Found ✅")
        urls = sitemap.get("sample_urls", [])
        if urls:
            doc.add_paragraph("Sample URLs:")
            for u in urls[:10]:
                p = doc.add_paragraph(u, style="List Bullet")
                p.runs[0].font.size = Pt(9)
    else:
        doc.add_paragraph("❌ Sitemap not found — generate and submit to Google Search Console")
        _code(doc, """# For Next.js — install next-sitemap package
# npm install next-sitemap

# next-sitemap.config.js
module.exports = {
  siteUrl: 'https://yoursite.com',
  generateRobotsTxt: true,
  exclude: ['/admin', '/user/*', '/api/*'],
}

# Then add to package.json scripts:
"postbuild": "next-sitemap"

# Submit at: https://search.google.com/search-console/sitemaps""", "Sitemap Generation Guide")

    _page_break(doc)

    # Dev checklist
    _h1(doc, "6. Developer Technical Checklist")
    checklist = [
        ("HTTPS redirect", "Ensure HTTP → HTTPS 301 redirect on all URLs"),
        ("Canonical tags", "Add <link rel='canonical'> to every page"),
        ("robots.txt", "Create /robots.txt — allow important pages, block admin/api"),
        ("sitemap.xml", "Generate and submit to Google Search Console"),
        ("Schema markup", "Add JSON-LD structured data to all key pages"),
        ("Image alt text", "Add descriptive alt text to every <img> tag"),
        ("Page speed", "Target Core Web Vitals: LCP < 2.5s, FID < 100ms, CLS < 0.1"),
        ("Mobile viewport", "Add <meta name='viewport' content='width=device-width'>"),
        ("Lazy loading", "Add loading='lazy' to all below-fold images"),
        ("Compression", "Enable Gzip/Brotli on server for HTML/CSS/JS"),
        ("Cache headers", "Set Cache-Control for static assets (max-age=31536000)"),
        ("Broken links", "Run crawl to find 404 errors and fix or redirect"),
        ("Hreflang", "Add hreflang if targeting multiple languages/regions"),
        ("JS rendering", "Enable SSR or SSG for React/Next.js pages (Google needs HTML)"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for c, t in zip(table.rows[0].cells, ["Task", "Description", "Done?"]):
        _set_cell_bg(c, "2C3E50")
        _cell_text(c, t, bold=True, color=C_WHITE, size=9)
    for task, desc in checklist:
        row = table.add_row().cells
        _set_cell_bg(row[0], "EAF0FB")
        _cell_text(row[0], task, bold=True, size=9, color=C_BLUE)
        _cell_text(row[1], desc, size=9)
        _cell_text(row[2], "☐", size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        for c in row:
            _set_cell_borders(c)

    _footer(doc, technical_data.get("base_url", ""))
    doc.save(output_path)


# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 2 — On-Page SEO Audit (Advanced)
# ══════════════════════════════════════════════════════════════════════════════

def create_pagewise_seo_doc(pages_analysis: list, output_path: str):
    doc = Document()
    _setup_margins(doc)

    # Cover
    _cover(doc, "On-Page SEO Audit Report", f"{len(pages_analysis)} Pages Analyzed",
           _avg_score(pages_analysis))

    # Executive summary
    _h1(doc, "Executive Summary")
    total_issues = sum(len(p.get("issues", [])) for p in pages_analysis)
    critical_count = sum(
        sum(1 for i in p.get("issues", []) if i.get("priority") == "critical")
        for p in pages_analysis
    )
    high_count = sum(
        sum(1 for i in p.get("issues", []) if i.get("priority") == "high")
        for p in pages_analysis
    )
    _kv(doc, "Total Pages Analyzed", len(pages_analysis))
    _kv(doc, "Total Issues Found", total_issues)
    _kv(doc, "Critical Issues", critical_count, highlight=True)
    _kv(doc, "High Priority Issues", high_count)
    _kv(doc, "Average SEO Score", f"{_avg_score(pages_analysis)}/100")

    # All crawled pages with URLs and summaries
    doc.add_paragraph()
    _h2(doc, "All Crawled Pages — URL & Summary")
    _pages_overview_table(doc, pages_analysis)

    # Score summary table
    doc.add_paragraph()
    _h2(doc, "Page Score Summary")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    for c, t in zip(table.rows[0].cells, ["Page", "Type", "Score", "Critical", "Issues"]):
        _set_cell_bg(c, "1A73E8")
        _cell_text(c, t, bold=True, color=C_WHITE, size=9)
    for p in pages_analysis:
        score = p.get("score", 0)
        crit = sum(1 for i in p.get("issues", []) if i.get("priority") == "critical")
        row = table.add_row().cells
        _set_cell_bg(row[0], "FAFAFA")
        _cell_text(row[0], p.get("url", "")[-50:], size=8)
        _cell_text(row[1], p.get("page_type", "").title(), size=9)
        score_color = C_LOW if score >= 80 else C_MEDIUM if score >= 50 else C_CRITICAL
        _cell_text(row[2], f"{score}/100", bold=True, color=score_color, size=10)
        _cell_text(row[3], str(crit) if crit else "✅ 0", bold=(crit > 0),
                   color=C_CRITICAL if crit > 0 else C_LOW, size=9)
        _cell_text(row[4], str(len(p.get("issues", []))), size=9)
        for c in row:
            _set_cell_borders(c)

    _page_break(doc)

    # ── Per-page detailed analysis ────────────────────────────────────────────
    for i, page in enumerate(pages_analysis, 1):
        url = page.get("url", "")
        page_type = page.get("page_type", "page").title()
        score = page.get("score", 0)
        scores = page.get("scores", {})
        issues = page.get("issues", [])
        current = page.get("current", {})
        sugg = page.get("suggestions", {})
        keywords = page.get("keywords", [])
        readability = page.get("readability", {})
        content_improvements = sugg.get("content_improvements", [])

        _h1(doc, f"Page {i} — {page_type}")
        _kv(doc, "URL", url)
        _score_badge(doc, "Overall Score", score)

        # Score breakdown
        _h2(doc, "Score Breakdown")
        score_table = doc.add_table(rows=1, cols=6)
        score_table.style = "Table Grid"
        for c, t in zip(score_table.rows[0].cells,
                        ["Title", "Meta Desc", "H1", "Content", "Technical", "Overall"]):
            _set_cell_bg(c, "34495E")
            _cell_text(c, t, bold=True, color=C_WHITE, size=9)
        row = score_table.add_row().cells
        for cell, key in zip(row, ["title", "meta", "h1", "content", "technical", "overall"]):
            val = scores.get(key, 0)
            clr = C_LOW if val >= 80 else C_MEDIUM if val >= 50 else C_CRITICAL
            _set_cell_bg(cell, "F8F9FA")
            _cell_text(cell, f"{val}/100", bold=True, color=clr, size=11,
                       align=WD_ALIGN_PARAGRAPH.CENTER)
            _set_cell_borders(cell)
        doc.add_paragraph()

        # Issues
        _h2(doc, f"Issues Found ({len(issues)})")
        _issues_table(doc, issues)
        doc.add_paragraph()

        # Before / After comparison
        _h2(doc, "Current vs. Suggested — Copy the Suggested Column")
        ba_rows = [
            ("Title Tag", current.get("title"), sugg.get("title")),
            ("Meta Description", current.get("meta_description"), sugg.get("meta_description")),
            ("H1 Tag", ", ".join(current.get("h1", [])) or None, sugg.get("h1")),
            ("URL Slug", current.get("canonical") or url, sugg.get("slug")),
            ("Canonical URL", current.get("canonical") or None, sugg.get("canonical")),
        ]
        _before_after_table(doc, ba_rows)
        doc.add_paragraph()

        # Stats row
        _h2(doc, "Page Statistics")
        stats = [
            ("Title Length", f"{current.get('title_length',0)} chars (ideal 50–60)"),
            ("Meta Length", f"{current.get('meta_length',0)} chars (ideal 150–160)"),
            ("H1 Count", f"{len(current.get('h1',[]))} (ideal = 1)"),
            ("H2 Count", f"{len(current.get('h2',[]))} (ideal ≥ 3)"),
            ("Word Count", f"{current.get('word_count',0)} words (min 300+)"),
            ("Internal Links", f"{current.get('internal_links_count',0)} (ideal 3–5+)"),
            ("Images no Alt", str(current.get("images_without_alt", 0))),
            ("Schema Markup", f"{current.get('schema_count',0)} found"),
            ("Render Method", page.get("render_method", "requests")),
            ("Readability", f"{readability.get('grade','—')} ({readability.get('avg_sentence_length',0)} avg words/sentence)"),
        ]
        stat_table = doc.add_table(rows=0, cols=2)
        stat_table.style = "Table Grid"
        for k, v in stats:
            row = stat_table.add_row().cells
            _set_cell_bg(row[0], "EAF0FB")
            _set_cell_bg(row[1], "FAFAFA")
            _cell_text(row[0], k, bold=True, size=9, color=C_BLUE)
            _cell_text(row[1], v, size=9)
            for c in row:
                _set_cell_borders(c)
        doc.add_paragraph()

        # Keywords
        if keywords:
            _h2(doc, "Top Keywords Detected on This Page")
            _keyword_table(doc, keywords)
            doc.add_paragraph()

        # OG Tags
        _h2(doc, "Open Graph (Social Sharing) Tags")
        og = current
        og_rows = [
            ("og:title", og.get("og_title"), sugg.get("title")),
            ("og:description", og.get("og_description"), sugg.get("meta_description")),
            ("og:image", og.get("og_image"), "[Upload 1200x630px image and add URL]"),
        ]
        _before_after_table(doc, og_rows)
        doc.add_paragraph()

        # Schema markup
        _h2(doc, "Schema JSON-LD (Structured Data)")
        schema = sugg.get("schema_json_ld")
        if schema:
            doc.add_paragraph(f"Recommended schema type: {schema.get('@type','WebPage')} — paste in <head>")
            _code(doc, json.dumps(schema, indent=2, ensure_ascii=False), "JSON-LD Schema")
        doc.add_paragraph()

        # Image alts
        image_alts = sugg.get("image_alts", [])
        if image_alts:
            _h2(doc, f"Images Missing Alt Text ({len(image_alts)} images)")
            img_table = doc.add_table(rows=1, cols=2)
            img_table.style = "Table Grid"
            for c, t in zip(img_table.rows[0].cells, ["Image Filename / URL", "Suggested Alt Text"]):
                _set_cell_bg(c, "E74C3C")
                _cell_text(c, t, bold=True, color=C_WHITE, size=9)
            for alt_item in image_alts[:10]:
                row = img_table.add_row().cells
                src = alt_item.get("src", "")[-60:]
                _set_cell_bg(row[0], "FFF5F5")
                _set_cell_bg(row[1], "F0FFF4")
                _cell_text(row[0], src, size=8, color=C_GRAY)
                _cell_text(row[1], alt_item.get("suggested_alt", ""), size=9, color=RGBColor(0x0D, 0x6E, 0x3B))
                for c in row:
                    _set_cell_borders(c)
            doc.add_paragraph()

        # Content improvements
        if content_improvements:
            _h2(doc, "Content Improvement Suggestions")
            for ci in content_improvements:
                _h3(doc, f"⚠ {ci.get('issue','')}")
                _kv(doc, "Fix", ci.get("fix",""))
                p = doc.add_paragraph()
                run = p.add_run(f"Example: {ci.get('example','')}")
                run.font.italic = True
                run.font.size = Pt(9)
                run.font.color.rgb = C_GRAY
            doc.add_paragraph()

        # Internal links
        int_links = sugg.get("internal_link_suggestions", [])
        if int_links:
            _h2(doc, "Internal Link Strategy")
            il_table = doc.add_table(rows=1, cols=3)
            il_table.style = "Table Grid"
            for c, t in zip(il_table.rows[0].cells, ["Anchor Text", "Link Target", "Why"]):
                _set_cell_bg(c, "1A73E8")
                _cell_text(c, t, bold=True, color=C_WHITE, size=9)
            for link in int_links[:6]:
                row = il_table.add_row().cells
                for c in row:
                    _set_cell_bg(c, "FAFAFA")
                    _set_cell_borders(c)
                _cell_text(row[0], link.get("anchor",""), size=9, color=C_BLUE)
                _cell_text(row[1], link.get("target",""), size=8, color=C_GRAY)
                _cell_text(row[2], link.get("reason",""), size=9)
            doc.add_paragraph()

        if i < len(pages_analysis):
            _page_break(doc)

    _footer(doc, pages_analysis[0].get("url","") if pages_analysis else "")
    doc.save(output_path)


# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 3 — Developer Implementation Guide
# ══════════════════════════════════════════════════════════════════════════════

def create_developer_guide_doc(pages_analysis: list, technical_data: dict, output_path: str):
    doc = Document()
    _setup_margins(doc)

    base_url = technical_data.get("base_url", "")
    _cover(doc, "Developer SEO Implementation Guide", base_url, None)

    doc.add_paragraph(
        "This guide contains copy-paste ready code snippets for every SEO fix. "
        "Review with a human before implementing. Do NOT auto-apply changes."
    )
    _divider(doc)

    # Global head template
    _h1(doc, "1. Global <head> Template (Every Page)")
    doc.add_paragraph("Add this in your base layout / _document.js / app.html:")
    _code(doc, """<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">

  <!-- SEO: Change per page -->
  <title>[Page Specific Title — 50-60 characters]</title>
  <meta name="description" content="[Page specific description — 150-160 chars]">
  <link rel="canonical" href="[Exact page URL]">
  <meta name="robots" content="index, follow">

  <!-- Open Graph (Social Sharing) -->
  <meta property="og:title" content="[Page Title]">
  <meta property="og:description" content="[Page Description]">
  <meta property="og:image" content="[Featured image 1200x630px]">
  <meta property="og:url" content="[Page URL]">
  <meta property="og:type" content="website">
  <meta property="og:site_name" content="[Brand Name]">

  <!-- Twitter Card -->
  <meta name="twitter:card" content="summary_large_image">
  <meta name="twitter:title" content="[Page Title]">
  <meta name="twitter:description" content="[Description]">
  <meta name="twitter:image" content="[Image URL]">

  <!-- Schema JSON-LD: Change per page type -->
  <script type="application/ld+json">
  { "@context": "https://schema.org", "@type": "WebPage", ... }
  </script>
</head>""", "HTML Template")

    _h2(doc, "For Next.js Projects — Use next-seo (Recommended)")
    _code(doc, """# Install
npm install next-seo

# _app.js — Default SEO for all pages
import { DefaultSeo } from 'next-seo';
export default function App({ Component, pageProps }) {
  return (
    <>
      <DefaultSeo
        titleTemplate="%s | YourBrand"
        defaultTitle="YourBrand — Sell & Buy Online"
        description="[Default site description]"
        openGraph={{ type: 'website', locale: 'en_IN', site_name: 'YourBrand' }}
      />
      <Component {...pageProps} />
    </>
  );
}

# Individual page — pages/sell-phone.js
import { NextSeo } from 'next-seo';
export default function SellPhone() {
  return (
    <>
      <NextSeo
        title="Sell Old Phone — Get Best Price Instantly"
        description="Sell your old mobile online and get instant cash. Free pickup, same-day payment."
        canonical="https://yoursite.com/sell-phone"
        openGraph={{
          url: 'https://yoursite.com/sell-phone',
          images: [{ url: 'https://yoursite.com/og-sell.jpg', width: 1200, height: 630 }],
        }}
      />
      <main>...</main>
    </>
  );
}""", "Next.js next-seo setup")

    _page_break(doc)

    # Per-page snippets
    _h1(doc, "2. Page-by-Page Implementation Snippets")
    for i, page in enumerate(pages_analysis[:8], 1):
        url = page.get("url", "")
        page_type = page.get("page_type", "page").title()
        hints = page.get("suggestions", {}).get("developer_hints", [])

        _h2(doc, f"2.{i} {page_type} Page")
        _kv(doc, "URL", url)
        doc.add_paragraph()

        for hint in hints:
            pri = hint.get("priority", "low")
            label = PRIORITY_LABEL.get(pri, pri)
            _h3(doc, f"{label} — {hint.get('type','')}")
            p = doc.add_paragraph()
            desc_run = p.add_run(f"Location: {hint.get('location','')}  |  {hint.get('description','')}")
            desc_run.font.size = Pt(9)
            desc_run.font.italic = True
            desc_run.font.color.rgb = C_GRAY
            _code(doc, hint.get("code",""), hint.get("type",""))
            doc.add_paragraph()

    _page_break(doc)

    # Image optimization
    _h1(doc, "3. Image SEO — Copy-Paste Templates")
    _code(doc, """<!-- Standard image with SEO -->
<img
  src="/images/sell-iphone-13.webp"
  alt="Sell iPhone 13 — Get Best Price Online"
  width="800"
  height="600"
  loading="lazy"
  decoding="async"
>

<!-- Responsive image with WebP + fallback -->
<picture>
  <source srcset="/images/product.webp" type="image/webp">
  <source srcset="/images/product.jpg" type="image/jpeg">
  <img src="/images/product.jpg" alt="[Descriptive alt text]" loading="lazy" width="800" height="600">
</picture>

<!-- Next.js Image component (automatic optimization) -->
import Image from 'next/image';
<Image
  src="/images/sell-phone.webp"
  alt="Sell your old phone — get instant cash"
  width={800}
  height={600}
  priority={false}  // true for hero images (above fold)
  placeholder="blur"
/>""", "Image HTML Templates")

    _page_break(doc)

    # Schema templates per page type
    _h1(doc, "4. Schema Markup Templates by Page Type")
    schemas = {
        "Homepage / Organization": {
            "@context": "https://schema.org",
            "@type": "Organization",
            "name": "[Brand Name]",
            "url": "[Homepage URL]",
            "logo": "[Logo URL — 600x60px]",
            "description": "[Brand description]",
            "contactPoint": {"@type": "ContactPoint", "telephone": "+91-XXXXXXXXXX", "contactType": "customer service"},
            "sameAs": ["[Facebook]", "[Instagram]", "[LinkedIn]"],
        },
        "Product Page": {
            "@context": "https://schema.org",
            "@type": "Product",
            "name": "[Product Name]",
            "description": "[Product description]",
            "image": "[Product Image URL]",
            "sku": "[SKU]",
            "brand": {"@type": "Brand", "name": "[Brand]"},
            "offers": {"@type": "Offer", "priceCurrency": "INR", "price": "[Price]",
                       "availability": "https://schema.org/InStock"},
            "aggregateRating": {"@type": "AggregateRating", "ratingValue": "4.5", "reviewCount": "48"},
        },
        "Category Page": {
            "@context": "https://schema.org",
            "@type": "CollectionPage",
            "name": "[Category Name]",
            "description": "[Category description]",
            "url": "[Category URL]",
            "breadcrumb": {"@type": "BreadcrumbList", "itemListElement": [
                {"@type": "ListItem", "position": 1, "name": "Home", "item": "[Home URL]"},
                {"@type": "ListItem", "position": 2, "name": "[Category]", "item": "[Category URL]"},
            ]},
        },
        "Blog Article": {
            "@context": "https://schema.org",
            "@type": "Article",
            "headline": "[Article Title]",
            "author": {"@type": "Person", "name": "[Author Name]"},
            "datePublished": "YYYY-MM-DD",
            "dateModified": "YYYY-MM-DD",
            "image": "[Featured Image URL]",
            "publisher": {"@type": "Organization", "name": "[Brand]",
                          "logo": {"@type": "ImageObject", "url": "[Logo URL]"}},
        },
    }
    for schema_name, schema_obj in schemas.items():
        _h2(doc, schema_name)
        _code(doc, f'<script type="application/ld+json">\n{json.dumps(schema_obj, indent=2)}\n</script>', "JSON-LD")
        doc.add_paragraph()

    _footer(doc, base_url)
    doc.save(output_path)


# ══════════════════════════════════════════════════════════════════════════════
# DOCUMENT 4 — GTM / GA4 / Meta Pixel Guide
# ══════════════════════════════════════════════════════════════════════════════

def create_tracking_guide_doc(tracking_data: dict, output_path: str):
    doc = Document()
    _setup_margins(doc)
    _cover(doc, "GTM / GA4 / Meta Pixel Guide", "Copy-paste tracking implementation", None)

    # GTM
    gtm = tracking_data.get("gtm", {})
    _h1(doc, "1. " + gtm.get("title", "Google Tag Manager"))
    for i, step in enumerate(gtm.get("steps", []), 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(step).font.size = Pt(10)

    _h2(doc, "Paste in <head>")
    _code(doc, gtm.get("head_snippet", ""), "HTML")
    _h2(doc, "Paste after <body>")
    _code(doc, gtm.get("body_snippet", ""), "HTML")
    p = doc.add_paragraph()
    p.add_run(f"⚠ {gtm.get('note','')}").font.color.rgb = C_HIGH

    _page_break(doc)

    # GA4
    ga4 = tracking_data.get("ga4", {})
    _h1(doc, "2. " + ga4.get("title", "Google Analytics 4"))
    for i, step in enumerate(ga4.get("steps", []), 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(step).font.size = Pt(10)
    _h2(doc, "GA4 Direct Code (without GTM)")
    _code(doc, ga4.get("direct_snippet", ""), "HTML")
    _h2(doc, "GA4 E-Commerce Events")
    for ev in ga4.get("ecommerce_events", []):
        _h3(doc, f"{ev['event']} — Trigger: {ev['trigger']}")
        _code(doc, ev["code"], "JavaScript")
    _page_break(doc)

    # Meta Pixel
    meta = tracking_data.get("meta_pixel", {})
    _h1(doc, "3. " + meta.get("title", "Meta Pixel"))
    for i, step in enumerate(meta.get("steps", []), 1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(step).font.size = Pt(10)
    _h2(doc, "Base Pixel Code — Paste in <head>")
    _code(doc, meta.get("base_code", ""), "HTML")
    _h2(doc, "Meta Pixel E-Commerce Events")
    for ev in meta.get("ecommerce_events", []):
        _h3(doc, f"{ev['event']} — {ev['trigger']}")
        _code(doc, ev["code"], "JavaScript")
    _page_break(doc)

    # DataLayer cheat sheet
    _h1(doc, "4. GTM DataLayer Cheat Sheet")
    _code(doc, """// ── Initialize dataLayer (put in <head> BEFORE GTM snippet) ──
window.dataLayer = window.dataLayer || [];

// ── Page view with custom data ──
dataLayer.push({
  event: 'page_view',
  page_type: 'product',        // homepage | category | product | cart | checkout
  page_title: document.title,
  page_url: window.location.href,
  user_logged_in: false,
});

// ── Product viewed ──
dataLayer.push({
  event: 'view_item',
  ecommerce: {
    items: [{
      item_id: 'SKU_001', item_name: 'iPhone 13 128GB',
      item_category: 'Smartphones', item_brand: 'Apple',
      price: 45000, quantity: 1
    }]
  }
});

// ── Add to cart ──
dataLayer.push({
  event: 'add_to_cart',
  ecommerce: {
    currency: 'INR', value: 45000,
    items: [{ item_id: 'SKU_001', item_name: 'iPhone 13', price: 45000, quantity: 1 }]
  }
});

// ── Purchase complete ──
dataLayer.push({
  event: 'purchase',
  ecommerce: {
    transaction_id: 'TXN_12345', value: 45000, currency: 'INR',
    items: [{ item_id: 'SKU_001', item_name: 'iPhone 13', price: 45000, quantity: 1 }]
  }
});

// ── Sell form submit (for your use case: sell old phone) ──
dataLayer.push({
  event: 'sell_form_submit',
  device_brand: 'Apple', device_model: 'iPhone 12',
  offered_price: 18000,
});""", "JavaScript DataLayer")

    _footer(doc, "")
    doc.save(output_path)


# ══════════════════════════════════════════════════════════════════════════════
# Helpers
# ══════════════════════════════════════════════════════════════════════════════

def _setup_margins(doc):
    for section in doc.sections:
        section.top_margin    = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin   = Cm(2.5)
        section.right_margin  = Cm(2.5)


def _avg_score(pages):
    scores = [p.get("score", 0) for p in pages]
    return int(sum(scores) / len(scores)) if scores else 0


def _cover(doc, title: str, subtitle: str, score):
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = C_BLUE

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(subtitle)
    r2.font.size = Pt(12)
    r2.font.color.rgb = C_GRAY

    if score is not None:
        p3 = doc.add_paragraph()
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = p3.add_run(f"Overall Score: {score}/100")
        r3.font.size = Pt(16)
        r3.font.bold = True
        r3.font.color.rgb = C_LOW if score >= 80 else C_MEDIUM if score >= 50 else C_CRITICAL

    p4 = doc.add_paragraph()
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r4 = p4.add_run(f"Generated by SEO AI Agent  |  {datetime.now().strftime('%d %B %Y')}")
    r4.font.size = Pt(9)
    r4.font.color.rgb = C_GRAY

    doc.add_paragraph()
    _divider(doc)
    doc.add_paragraph()


# ══════════════════════════════════════════════════════════════════════════════
# Main entry point
# ══════════════════════════════════════════════════════════════════════════════

def generate_all_documents(audit_data: dict, output_dir: str) -> str:
    os.makedirs(output_dir, exist_ok=True)

    technical    = audit_data.get("technical_seo", {})
    pages        = audit_data.get("pages_analysis", [])
    tracking     = audit_data.get("tracking_guide", {})

    paths = {
        "technical": os.path.join(output_dir, "1_Technical_SEO_Audit.docx"),
        "pagewise":  os.path.join(output_dir, "2_OnPage_SEO_Audit.docx"),
        "developer": os.path.join(output_dir, "3_Developer_Implementation_Guide.docx"),
        "tracking":  os.path.join(output_dir, "4_GTM_GA4_MetaPixel_Guide.docx"),
    }

    create_technical_seo_doc(technical, paths["technical"], pages_analysis=pages)
    if pages:
        create_pagewise_seo_doc(pages, paths["pagewise"])
    create_developer_guide_doc(pages, technical, paths["developer"])
    create_tracking_guide_doc(tracking, paths["tracking"])

    zip_path = os.path.join(output_dir, "SEO_Audit_Reports.zip")
    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        for path in paths.values():
            if os.path.exists(path):
                zf.write(path, os.path.basename(path))

    return zip_path
